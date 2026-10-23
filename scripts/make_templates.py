"""Builds templates/cv_template.docx and templates/cl_template.docx.

The CV template recreates the design of the original one-page CV (Letter, Times, grey-shaded entry headers,
❖ bullets, right-aligned dates). The {{ }} / {%p %} tags are docxtpl slots. Only text goes into them,
so the design can't change. You can also open the .docx in Word and restyle it by hand. Keep the tags intact.

Run:  .venv/bin/python scripts/make_templates.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FONT = "Times New Roman"
BODY_PT = 10
TEXT_WIDTH = Inches(7.5)  # Letter 8.5in minus 2 x 0.5in margins
BULLET_INDENT = Inches(0.22)
SHADE = "D9D9D9"


def base_document(margins_in: float, top_in: float) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(margins_in)
    sec.top_margin = sec.bottom_margin = Inches(top_in)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return doc


def para(doc, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    return p


def run(p, text, bold=False, italic=False, size=None, shade=False, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if shade:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), SHADE)
        r._r.get_or_add_rPr().append(shd)
    return r


def right_tab(p):
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)


def tag(doc, text):
    """A {%p ... %} control paragraph (removed by docxtpl at render time)."""
    doc.add_paragraph(text)


def bottom_border(p):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def section_heading(doc, text):
    p = para(doc)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    run(p, text, bold=True, size=11)
    bottom_border(p)


def entry_header(doc, left_runs, right_text, right_bold=False, right_italic=True):
    p = para(doc)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.keep_with_next = True
    right_tab(p)
    for text, bold in left_runs:
        run(p, text, bold=bold, shade=True)
    run(p, "\t")
    run(p, right_text, bold=right_bold, italic=right_italic)
    return p


def bullet(doc, runs, right=None):
    """runs: list of (text, bold, italic). The ❖ glyph and hanging indent imitate the original CV."""
    p = para(doc)
    pf = p.paragraph_format
    pf.left_indent = BULLET_INDENT
    pf.first_line_indent = -BULLET_INDENT + Inches(0.05)
    pf.tab_stops.add_tab_stop(BULLET_INDENT, WD_TAB_ALIGNMENT.LEFT)
    if right:
        right_tab(p)
    run(p, "❖", size=7)
    run(p, "\t")
    for text, bold, italic in runs:
        run(p, text, bold=bold, italic=italic)
    if right:
        run(p, "\t")
        run(p, right[0], bold=right[1], italic=right[2])
    return p


def build_cv(path: Path):
    doc = base_document(0.5, 0.4)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "{{ personal.name|upper }}", bold=True, size=17)
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "{{ personal.contact_line }}", italic=True, size=9.5)

    section_heading(doc, "EDUCATION")
    tag(doc, "{%p for e in education %}")
    entry_header(doc, [("{{ e.institution }}{% if e.location %} | {{ e.location }}{% endif %}", True)], "{{ e.dates }}")
    bullet(doc, [("{{ e.degree }}", True, False)], right=("{{ e.gpa }}", True, True))
    tag(doc, "{%p for d in e.details %}")
    bullet(doc, [("{{ d }}", False, False)])
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if e.coursework %}")
    p = para(doc)
    run(p, "Relevant Coursework: ", bold=True)
    run(p, "{{ e.coursework }}")
    tag(doc, "{%p endif %}")
    tag(doc, "{%p endfor %}")

    section_heading(doc, "EXPERIENCE")
    tag(doc, "{%p for job in experience %}")
    entry_header(doc, [("{{ job.company }} | {{ job.title }} | {{ job.location }}", True)], "{{ job.dates }}")
    tag(doc, "{%p for b in job.bullets %}")
    p = bullet(doc, [])
    run(p, "{{r b }}")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endfor %}")

    section_heading(doc, "PROJECTS")
    tag(doc, "{%p for proj in projects %}")
    entry_header(doc, [("{{ proj.name }}", True)], "{{ proj.stack }}", right_bold=True, right_italic=True)
    tag(doc, "{%p for b in proj.bullets %}")
    p = bullet(doc, [])
    run(p, "{{r b }}")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endfor %}")

    section_heading(doc, "SKILLS & INTERESTS")
    tag(doc, "{%p for s in skills %}")
    bullet(doc, [("{{ s.category }}:", True, False), (" {{ s.list }}", False, False)])
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if languages_line %}")
    bullet(doc, [("{{ languages_label }}:", True, False), (" {{ languages_line }}", False, False)])
    tag(doc, "{%p endif %}")

    doc.save(path)


def build_cl(path: Path):
    doc = base_document(1.0, 1.0)
    for style in (doc.styles["Normal"],):
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(0)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "{{ personal.name|upper }}", bold=True, size=16)
    bottom_border(p)

    para(doc)
    para(doc).add_run("{{ date }}")
    para(doc)
    tag(doc, "{%p for line in recipient_lines %}")
    para(doc).add_run("{{ line }}")
    tag(doc, "{%p endfor %}")
    para(doc)
    p = para(doc)
    p.paragraph_format.space_after = Pt(8)
    p.add_run("{{ salutation }}")
    tag(doc, "{%p for para in paragraphs %}")
    p = para(doc)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run("{{ para }}")
    tag(doc, "{%p endfor %}")
    para(doc).add_run("{{ closing }}")
    para(doc).add_run("{{ personal.name }}")
    p = para(doc)
    run(p, "{{ personal.phone }} | {{ personal.email }} | {{ personal.github }}", size=10)

    doc.save(path)


if __name__ == "__main__":
    out = ROOT / "templates"
    out.mkdir(exist_ok=True)
    build_cv(out / "cv_template.docx")
    build_cl(out / "cl_template.docx")
    print("Wrote", out / "cv_template.docx", "and", out / "cl_template.docx")
