"""Bootstrap profile.yaml from an existing CV (.docx or .pdf). You review the result before saving."""
from __future__ import annotations

import io
import re

import yaml
from docx import Document
from pypdf import PdfReader

from . import prompts
from .llm import chat_json

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\w)\+?\(?\d[\d\s().-]{7,}\d(?!\w)")


def redact_contacts(text: str) -> tuple[str, dict]:
    """Remove email and phone before the CV text leaves this machine. Returns the text and what was removed."""
    found = {}

    def phone(m: re.Match) -> str:
        if sum(c.isdigit() for c in m.group()) < 10:  # date ranges like 2022-2024, not a phone number
            return m.group()
        found.setdefault("phone", m.group().strip())
        return "[phone removed]"

    if m := EMAIL_RE.search(text):
        found["email"] = m.group()
    text = EMAIL_RE.sub("[email removed]", text)
    return PHONE_RE.sub(phone, text), found


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(p for p in parts if p.strip())
    if name.endswith(".pdf"):
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    return data.decode("utf-8", errors="ignore")


def draft_profile_yaml(cv_text: str, model: str) -> str:
    cv_text, contacts = redact_contacts(cv_text)
    data = chat_json(prompts.IMPORT_SYSTEM, prompts.IMPORT_USER.format(text=cv_text), model, temperature=0, task="import CV")
    personal = data.setdefault("personal", {}) if isinstance(data.get("personal", {}), dict) else {}
    for key in ("email", "phone"):
        if contacts.get(key) or "removed" in str(personal.get(key, "")):
            personal[key] = contacts.get(key, "")
    header = (
        "# Drafted from your CV by the importer. REVIEW EVERY LINE before saving.\n"
        "# Add more facts freely — the more true detail here, the better the tailoring.\n"
        "# Facts can be plain strings or {text: ..., skills: [...]}. Optional: id, max_bullets per job.\n\n"
    )
    return header + yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=120)
